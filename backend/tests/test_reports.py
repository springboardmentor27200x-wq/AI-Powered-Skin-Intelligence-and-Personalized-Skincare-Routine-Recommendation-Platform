from io import BytesIO

from openpyxl import load_workbook


def test_progress_reports_require_authentication(client):
    response = client.get("/api/reports/progress?format=pdf")
    assert response.status_code == 401


def test_progress_report_exports_pdf_and_xlsx(client, auth_headers):
    generated = client.post(
        "/api/routines/generate",
        headers=auth_headers,
        json={"age": 29, "skin_type": "combination", "concerns": ["redness"]},
    )
    assert generated.status_code == 200

    pdf = client.get("/api/reports/progress?format=pdf", headers=auth_headers)
    assert pdf.status_code == 200
    assert pdf.headers["content-type"].startswith("application/pdf")
    assert pdf.content.startswith(b"%PDF-")
    assert "skin-progress-report.pdf" in pdf.headers["content-disposition"]

    xlsx = client.get("/api/reports/progress?format=xlsx", headers=auth_headers)
    assert xlsx.status_code == 200
    assert xlsx.content.startswith(b"PK")
    workbook = load_workbook(BytesIO(xlsx.content), read_only=True)
    assert workbook.sheetnames == ["Progress Summary", "Assessment History", "Active Routine"]
    assert workbook["Assessment History"].max_row == 2


def test_progress_report_rejects_unsupported_format(client, auth_headers):
    response = client.get("/api/reports/progress?format=csv", headers=auth_headers)
    assert response.status_code == 422