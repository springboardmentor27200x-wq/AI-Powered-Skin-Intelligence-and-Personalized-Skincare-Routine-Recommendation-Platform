import io
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.system_report_engine import build_system_pdf, build_system_excel

STATS = SimpleNamespace(
    total_users=6, active_users=5, inactive_users=1, total_skin_profiles=3, total_daily_logs=9,
    users_by_role={"user": 2, "skincare_consultant": 1, "dermatologist": 1, "administrator": 2},
)
ANALYTICS = SimpleNamespace(
    total_assessments=62, average_score=81.0, total_routines=2, total_ingredients=8, total_products=16,
    top_concern="acne", top_concern_count=1, platform_adherence_rate=100.0,
)


def product(name="Salicylic Gel", category="Face Wash"):
    return SimpleNamespace(name=name, category=category, brand="DermaBasics", times_recommended=2,
                           average_score=92.5, allergen_conflicts=1, caution_flags=0)


def monitoring(**over):
    base = dict(
        users_evaluated=2, products_evaluated=16, recommended_min_score=80, total_pairs=32, average_score=88.4,
        allergen_conflict_pairs=3, conflict_rate=9.4,
        top_products=[product()], most_flagged_products=[product("Retinol Serum", "Serum")],
        category_scores=[SimpleNamespace(category="Face Wash", average_score=92.5, products=3)],
    )
    base.update(over)
    return SimpleNamespace(**base)


EMPTY = monitoring(users_evaluated=0, products_evaluated=0, total_pairs=0, average_score=None,
                   allergen_conflict_pairs=0, conflict_rate=None, top_products=[], most_flagged_products=[],
                   category_scores=[])


def cells(ws):
    return {(row[0].value): row[1].value for row in ws.iter_rows() if row[0].value is not None and len(row) > 1}


# ---- PDF ----

def test_pdf_is_a_valid_pdf():
    pdf = build_system_pdf("Admin", STATS, ANALYTICS, monitoring())
    assert pdf.startswith(b"%PDF") and len(pdf) > 1000


def test_pdf_builds_when_there_is_no_recommendation_data():
    assert build_system_pdf("Admin", STATS, ANALYTICS, EMPTY).startswith(b"%PDF")


def test_pdf_survives_special_characters():
    pdf = build_system_pdf("A & B <Admin>", STATS, ANALYTICS, monitoring(top_products=[product("Vitamin C & E <Serum>")]))
    assert pdf.startswith(b"%PDF")


def test_pdf_handles_missing_optional_numbers():
    analytics = SimpleNamespace(**{**ANALYTICS.__dict__, "average_score": None, "top_concern": None,
                                   "top_concern_count": 0, "platform_adherence_rate": None})
    assert build_system_pdf("Admin", STATS, analytics, monitoring()).startswith(b"%PDF")


# ---- Excel ----

def load(xlsx_bytes):
    return load_workbook(io.BytesIO(xlsx_bytes))


def test_excel_has_the_expected_sheets():
    wb = load(build_system_excel(STATS, ANALYTICS, monitoring()))
    assert wb.sheetnames == ["Summary", "Top products", "Flagged products", "Category scores"]


def test_excel_summary_values_are_real_numbers():
    ws = load(build_system_excel(STATS, ANALYTICS, monitoring()))["Summary"]
    values = cells(ws)
    assert values["Total users"] == 6 and isinstance(values["Total users"], int)
    assert values["Average skin score"] == 81.0
    assert values["Assessments run"] == 62
    assert values["Top concern"] == "Acne (1 user)"
    assert values["7-day routine adherence"] == "100.0%"


def test_excel_summary_includes_users_by_role_and_monitoring():
    values = cells(load(build_system_excel(STATS, ANALYTICS, monitoring()))["Summary"])
    assert values["Skincare Consultant"] == 1
    assert values["Users evaluated"] == 2
    assert values["Allergen conflict rate"] == "9.4%"


def test_excel_product_sheets_list_rows():
    wb = load(build_system_excel(STATS, ANALYTICS, monitoring()))
    top = list(wb["Top products"].iter_rows(values_only=True))
    assert top[0] == ("Product", "Category", "Brand", "Recommended for (users)", "Average score")
    assert top[1] == ("Salicylic Gel", "Face Wash", "DermaBasics", 2, 92.5)
    flagged = list(wb["Flagged products"].iter_rows(values_only=True))
    assert flagged[1][0] == "Retinol Serum" and flagged[1][2] == 1


def test_excel_with_no_data_has_headers_only():
    wb = load(build_system_excel(STATS, ANALYTICS, EMPTY))
    assert len(list(wb["Top products"].iter_rows())) == 1
    assert len(list(wb["Category scores"].iter_rows())) == 1
    assert cells(wb["Summary"])["Average recommendation score"] == "—"