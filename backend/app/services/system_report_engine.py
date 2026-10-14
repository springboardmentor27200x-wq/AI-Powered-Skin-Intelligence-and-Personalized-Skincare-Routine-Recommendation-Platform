"""Platform-wide system report for administrators (PDF and Excel).

Aggregates only — no individual user's personal data. It presents the same numbers the
admin dashboard shows: user stats, platform analytics, and recommendation monitoring.
"""
import io
from datetime import datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle

from app.services.report_engine import _pdf_styles, _pdf_header

TITLE = "Platform System Report"


def _label(word: str | None) -> str:
    return word.replace("_", " ").title() if word else "—"


def _plural(n: int, word: str) -> str:
    return f"{n} {word}{'' if n == 1 else 's'}"


def _top_concern_text(analytics) -> str:
    if not analytics.top_concern:
        return "—"
    return f"{_label(analytics.top_concern)} ({_plural(analytics.top_concern_count, 'user')})"


def _adherence_text(analytics) -> str:
    return f"{analytics.platform_adherence_rate}%" if analytics.platform_adherence_rate is not None else "No data yet"


def _user_rows(stats) -> list[list[str]]:
    rows = [
        ["Total users", str(stats.total_users)],
        ["Active", str(stats.active_users)],
        ["Deactivated", str(stats.inactive_users)],
    ]
    rows += [[_label(role), str(count)] for role, count in stats.users_by_role.items()]
    return rows


def _analytics_rows(stats, analytics) -> list[list[str]]:
    return [
        ["Skin profiles", str(stats.total_skin_profiles)],
        ["Daily logs", str(stats.total_daily_logs)],
        ["Assessments run", str(analytics.total_assessments)],
        ["Average skin score", str(analytics.average_score) if analytics.average_score is not None else "—"],
        ["Routines generated", str(analytics.total_routines)],
        ["Ingredients tracked", str(analytics.total_ingredients)],
        ["Products catalogued", str(analytics.total_products)],
        ["Top concern", _top_concern_text(analytics)],
        ["7-day routine adherence", _adherence_text(analytics)],
    ]


def _monitoring_rows(m) -> list[list[str]]:
    return [
        ["Users evaluated", str(m.users_evaluated)],
        ["Products evaluated", str(m.products_evaluated)],
        ["User–product pairings scored", str(m.total_pairs)],
        ["Average recommendation score", str(m.average_score) if m.average_score is not None else "—"],
        [f"\"Recommended\" threshold (score ≥)", str(m.recommended_min_score)],
        ["Pairings with an allergen conflict", str(m.allergen_conflict_pairs)],
        ["Allergen conflict rate", f"{m.conflict_rate}%" if m.conflict_rate is not None else "—"],
    ]


def _excel_value(value):
    """Summary rows are built as text for the PDF; give Excel real numbers where it can."""
    if isinstance(value, str):
        for cast in (int, float):
            try:
                return cast(value)
            except ValueError:
                pass
    return value


def _pdf_table(rows, col_widths):
    table = Table(rows, colWidths=col_widths)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF5F1")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#25543F")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D8CFC0")),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def build_system_pdf(admin_name: str, stats, analytics, monitoring) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = _pdf_styles()
    story = []
    _pdf_header(story, styles, TITLE, escape(admin_name))

    story.append(Paragraph("Users", styles["SectionHeading"]))
    story.append(_pdf_table([["Metric", "Value"]] + _user_rows(stats), [3.5 * inch, 2 * inch]))

    story.append(Paragraph("Platform analytics", styles["SectionHeading"]))
    story.append(_pdf_table([["Metric", "Value"]] + _analytics_rows(stats, analytics), [3.5 * inch, 2 * inch]))

    story.append(Paragraph("Recommendation monitoring", styles["SectionHeading"]))
    if monitoring.users_evaluated == 0:
        story.append(Paragraph(
            "No users have a skin profile yet, so there are no recommendations to monitor.", styles["Normal"]
        ))
    else:
        story.append(_pdf_table([["Metric", "Value"]] + _monitoring_rows(monitoring), [3.5 * inch, 2 * inch]))

        story.append(Paragraph("Most recommended products", styles["SectionHeading"]))
        if monitoring.top_products:
            rows = [["Product", "Category", "Recommended for", "Avg score"]]
            rows += [
                [p.name, p.category, _plural(p.times_recommended, "user"), str(p.average_score)]
                for p in monitoring.top_products
            ]
            story.append(_pdf_table(rows, [2.4 * inch, 1.4 * inch, 1.4 * inch, 0.9 * inch]))
        else:
            story.append(Paragraph("No product reaches the recommended threshold for any user.", styles["Normal"]))

        story.append(Paragraph("Most flagged products", styles["SectionHeading"]))
        if monitoring.most_flagged_products:
            rows = [["Product", "Category", "Allergen conflicts", "Cautions"]]
            rows += [
                [p.name, p.category, str(p.allergen_conflicts), str(p.caution_flags)]
                for p in monitoring.most_flagged_products
            ]
            story.append(_pdf_table(rows, [2.4 * inch, 1.4 * inch, 1.4 * inch, 0.9 * inch]))
        else:
            story.append(Paragraph("No product is flagged for any user.", styles["Normal"]))

        story.append(Paragraph("Average score by category", styles["SectionHeading"]))
        rows = [["Category", "Avg score", "Products"]]
        rows += [[c.category, str(c.average_score), str(c.products)] for c in monitoring.category_scores]
        story.append(_pdf_table(rows, [2.4 * inch, 1.4 * inch, 1.4 * inch]))

    doc.build(story)
    return buf.getvalue()


def build_system_excel(stats, analytics, monitoring) -> bytes:
    bold = Font(bold=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append([TITLE])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"])

    def block(title, rows):
        ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = bold
        for label, value in rows:
            ws.append([label, _excel_value(value)])

    block("Users", _user_rows(stats))
    block("Platform analytics", _analytics_rows(stats, analytics))
    block("Recommendation monitoring", _monitoring_rows(monitoring))
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 24

    def sheet(title, header, rows, widths):
        s = wb.create_sheet(title)
        s.append(header)
        for cell in s[1]:
            cell.font = bold
        for row in rows:
            s.append(row)
        for letter_, width in zip("ABCDE", widths):
            s.column_dimensions[letter_].width = width

    sheet(
        "Top products",
        ["Product", "Category", "Brand", "Recommended for (users)", "Average score"],
        [[p.name, p.category, p.brand or "", p.times_recommended, p.average_score] for p in monitoring.top_products],
        [30, 18, 18, 24, 15],
    )
    sheet(
        "Flagged products",
        ["Product", "Category", "Allergen conflicts", "Cautions", "Average score"],
        [[p.name, p.category, p.allergen_conflicts, p.caution_flags, p.average_score] for p in monitoring.most_flagged_products],
        [30, 18, 20, 12, 15],
    )
    sheet(
        "Category scores",
        ["Category", "Average score", "Products"],
        [[c.category, c.average_score, c.products] for c in monitoring.category_scores],
        [24, 16, 12],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()