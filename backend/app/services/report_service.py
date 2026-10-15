from datetime import datetime, timezone
from copy import copy
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.models.user import User
from app.schemas.dashboard import ProgressAnalyticsResponse
from app.services.dashboard_service import dashboard_service
from app.services.routine_service import get_current_routine


def _report_data(db: Session, user: User) -> tuple[ProgressAnalyticsResponse, object]:
    return dashboard_service.get_progress_analytics(db, user), get_current_routine(db, user)


def _weekly_adherence(routine) -> float:
    active_days = [day for day in routine.recent_progress if day.total_steps > 0]
    if not active_days:
        return 0.0
    return sum(day.adherence_percentage for day in active_days) / len(active_days)


def _build_pdf(user: User, progress: ProgressAnalyticsResponse, routine) -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#065f46")))
    story = [
        Paragraph("Skin Progress Report", styles["ReportTitle"]),
        Paragraph(f"Prepared for {escape(user.full_name)} · {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]),
        Spacer(1, 18),
        Paragraph("Current Scores", styles["Heading2"]),
    ]
    score_rows = [["Measure", "Score / 100"],
        ["Overall skin health", progress.overall_skin_health_score],
        ["Skin condition", progress.skin_condition_score],
        ["Lifestyle", progress.lifestyle_score],
        ["Sleep", progress.sleep_score],
        ["Hydration", progress.hydration_score],
    ]
    score_table = Table(score_rows, colWidths=[4.8 * inch, 1.6 * inch], repeatRows=1)
    score_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#065f46")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("PADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([score_table, Spacer(1, 14), Paragraph("Assessment History", styles["Heading2"])])
    history_rows = [["Saved", "Assessment score / 100"]]
    history_rows.extend([[point.label, point.score] for point in progress.trend])
    if len(history_rows) == 1:
        history_rows.append(["No saved assessment history", ""])
    history_table = Table(history_rows, colWidths=[4.8 * inch, 1.6 * inch], repeatRows=1)
    history_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("PADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([history_table, Spacer(1, 14)])
    if routine:
        story.append(Paragraph("Active Routine Progress", styles["Heading2"]))
        story.append(Paragraph(
            f"Today: {routine.today_completed_steps} of {routine.today_total_steps} morning/evening steps complete.",
            styles["Normal"],
        ))
        story.append(Paragraph(f"Seven-day average adherence: {_weekly_adherence(routine):.1f}%", styles["Normal"]))
        story.append(Spacer(1, 8))
    story.append(Paragraph(escape(progress.summary), styles["Normal"]))
    story.append(Spacer(1, 10))
    story.append(Paragraph(
        "This educational report is not a medical diagnosis. Consult a dermatologist about persistent or worsening symptoms.",
        styles["Italic"],
    ))
    document.build(story)
    return buffer.getvalue()


def _build_xlsx(user: User, progress: ProgressAnalyticsResponse, routine) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Progress Summary"
    summary.append(["Skin Progress Report"])
    summary.append(["Prepared for", user.full_name])
    summary.append(["Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
    summary.append([])
    summary.append(["Measure", "Score / 100"])
    for row in [
        ("Overall skin health", progress.overall_skin_health_score),
        ("Skin condition", progress.skin_condition_score),
        ("Lifestyle", progress.lifestyle_score),
        ("Sleep", progress.sleep_score),
        ("Hydration", progress.hydration_score),
    ]:
        summary.append(row)
    summary.append([])
    summary.append(["Summary", progress.summary])
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 76

    history = workbook.create_sheet("Assessment History")
    history.append(["Saved", "Assessment score / 100"])
    for point in progress.trend:
        history.append([point.label, point.score])
    if not progress.trend:
        history.append(["No saved assessment history", None])
    history.column_dimensions["A"].width = 24
    history.column_dimensions["B"].width = 26

    if routine:
        routine_sheet = workbook.create_sheet("Active Routine")
        routine_sheet.append(["Routine section", "Step", "Completed today"])
        completed = set(routine.completed_step_keys)
        for section in ("morning", "evening", "weekly"):
            for index, step in enumerate(routine.routine.get(section, [])):
                key = f"{section}:{index}"
                routine_sheet.append([section.title(), step, "Yes" if key in completed else "No" if section != "weekly" else ""])
        routine_sheet.append([])
        routine_sheet.append(["Seven-day average adherence (%)", round(_weekly_adherence(routine), 1)])
        routine_sheet.column_dimensions["A"].width = 24
        routine_sheet.column_dimensions["B"].width = 82
        routine_sheet.column_dimensions["C"].width = 22

    header_fill = PatternFill("solid", fgColor="065F46")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = copy(cell.alignment) + Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_progress_report(db: Session, user: User, file_format: str) -> tuple[bytes, str, str]:
    progress, routine = _report_data(db, user)
    if file_format == "pdf":
        return _build_pdf(user, progress, routine), "application/pdf", "skin-progress-report.pdf"
    return (
        _build_xlsx(user, progress, routine),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "skin-progress-report.xlsx",
    )